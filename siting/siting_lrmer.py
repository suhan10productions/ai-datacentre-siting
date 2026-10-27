"""Long-run marginal emission rates (LRMER) for the US AI campuses (Sections 4.2, 5.2).

Source: NREL Cambium 2023 LRMER workbook (Gagnon, 2024; data.nrel.gov/submissions/230,
file Cambium23_LRMER_GEAregions.xlsx), CO2 from direct combustion, kg per MWh of
end-use demand, by Grid and Emission Assessment (GEA) region, for 2025-2050 in
five-year steps. Each campus is placed in its county (plotly/datasets
geojson-counties-fips.json, US Census boundaries) and the county is mapped to its
GEA region with the workbook's County Mapping tab.

Levelisation follows the workbook's method: annual values are interpolated
linearly between the five-year points and averaged over 2027-2046 (a 20-year
life from a typical start in 2027) with a 3% real discount rate. The central
case is the Mid-case scenario; the range uses the Low and High Renewable Energy
Cost, High Demand Growth and Low and High Natural Gas Price scenarios.

US reference: the LRMER where the US urban population lives (Natural Earth
populated places in the contiguous US, weighted by pop_max), on the same logic
as the water benchmark (demand follows population).

Outputs results/results_lrmer.json.
"""
import json
import sys
import numpy as np
import pandas as pd
import geopandas as gpd
import openpyxl

EXT = "data/ext/"
START, LIFE, DISC = 2027, 20, 0.03
SCEN = ["Mid-case", "Low Renewable Energy Costs", "High Renewable Energy Costs",
        "Low Natural Gas Prices", "High Natural Gas Prices", "High Demand Growth"]
RNG = np.random.default_rng(20261027)
NBOOT = 5000

# ---- Cambium annual LRMER, CO2 combustion block
wb = openpyxl.load_workbook(EXT + "Cambium23_LRMER_GEAregions.xlsx", read_only=True, data_only=True)
rows = list(wb["Data - Annual"].iter_rows(values_only=True))
hdr_scen, hdr_year = rows[3], rows[4]
block = 15  # first column of "CO2 from Direct Combustion"
cols = {}
for j in range(block, block + 48):
    s = hdr_scen[j] if hdr_scen[j] else s
    cols[(s, int(hdr_year[j]))] = j
regions = [r[1] for r in rows[5:23]]
years = np.arange(2025, 2051, 5)
w_years = np.arange(START, START + LIFE)
weights = (1 + DISC) ** -(w_years - START)


def levelised(region_row, scen):
    pts = np.array([float(rows[region_row][cols[(scen, y)]]) for y in years])
    annual = np.interp(np.minimum(w_years, 2050), years, pts)
    return float((annual * weights).sum() / weights.sum())


LR = {s: {regions[i]: levelised(5 + i, s) for i in range(len(regions))} for s in SCEN}

# ---- county -> GEA
cm = list(wb["County Mapping"].iter_rows(values_only=True))
county_gea = {f"{int(r[0]):02d}{int(r[1]):03d}": r[6] for r in cm[1:] if r[0] is not None and r[6]}
counties = gpd.read_file(EXT + "counties.json")
counties["fips"] = counties["id"].astype(str).str.zfill(5)
counties["gea"] = counties.fips.map(county_gea)

def to_gea(lon, lat):
    pts = gpd.GeoDataFrame(geometry=gpd.points_from_xy(lon, lat), crs="EPSG:4326")
    j = gpd.sjoin(pts, counties[["gea", "geometry"]].set_crs("EPSG:4326", allow_override=True),
                  how="left", predicate="within")
    j = j[~j.index.duplicated()]
    return j["gea"].to_numpy()

# ---- campuses
import re
norm = lambda x: re.sub(r"[^a-z0-9]", "", str(x).lower())
e = pd.read_csv("results/results_epoch_sites.csv")
co = pd.read_csv("data/epoch_coords.csv")
co["k"] = co.site_id.map(norm); e["k"] = e.Name.map(norm)
left = sorted(set(e.k) - set(co.k))  # one coordinate entry has no slug (as in siting_analysis.py)
if (co.site_id == "UNKNOWN_New_Mexico").any() and len(left) == 1:
    co.loc[co.site_id == "UNKNOWN_New_Mexico", "k"] = left[0]
e = e.merge(co[["k", "lon", "lat"]], on="k", how="left")
assert e.lon.notna().all()
us = e[(e.name_0 == "United States") & (e.planned_it_mw > 0)].copy()
us["gea"] = to_gea(us.lon, us.lat)
assert us.gea.notna().all(), us[us.gea.isna()][["Name", "name_1"]]

# ---- population reference (contiguous US)
pl = gpd.read_file(EXT + "ne_places.geojson")
pl = pl[(pl.adm0name == "United States of America") & pl.longitude.between(-125, -66) & pl.latitude.between(24, 50)]
pl["gea"] = to_gea(pl.longitude, pl.latitude)
pl = pl[pl.gea.notna()]

st = pd.read_csv("data/ember_us_states_2024.psv", sep="|")
US_AVG = float(st[st.state_code == "US"].intensity_g_per_kwh.iloc[0])
w = us.planned_it_mw.to_numpy(float)
hw = (us.bws_cat >= 3).to_numpy()
ok_avg = ~(hw | (us.grid_ci.to_numpy() > US_AVG))
sys.path.insert(0, "../model"); sys.path.insert(0, "../work")
import model as mdl  # central facility energy per GW, as in siting_screen.py
E = mdl.RHO["central"] * mdl.load_factor(mdl.U["central"], mdl.IDLE["central"]) * mdl.HOURS / 1000 * mdl.PUE["central"]


def stats(scen, idx=None, pidx=None):
    m = us.gea.map(LR[scen]).to_numpy(float)
    pm = pl.gea.map(LR[scen]).to_numpy(float); pw = pl.pop_max.to_numpy(float)
    if idx is None: idx = np.arange(len(w))
    if pidx is None: pidx = np.arange(len(pm))
    ww, mm, oa, hh = w[idx], m[idx], ok_avg[idx], hw[idx]
    ref = (pm[pidx] * pw[pidx]).sum() / pw[pidx].sum()
    camp = (mm * ww).sum() / ww.sum()
    om = ~(hh | (mm > ref))
    f = ww.sum() / 1000 * E / 1000  # GW * TWh/GW -> TWh ... kg/MWh * TWh = kt; /1000 -> Mt
    pa = (mm[oa] * ww[oa]).sum() / ww[oa].sum()
    pmv = (mm[om] * ww[om]).sum() / ww[om].sum() if ww[om].sum() > 0 else np.nan
    return dict(ref=ref, campus=camp, campus_over_ref=camp / ref - 1,
                avgscreen_change_mt=-(f * (camp - pa)), lrscreen_pass=ww[om].sum() / ww.sum(),
                lrscreen_change_mt=-(f * (camp - pmv)))


out = dict(settings=dict(start=START, life=LIFE, discount=DISC, stage="combustion CO2", units="kg/MWh end-use"),
           levelised={s: LR[s] for s in SCEN}, n_places=int(len(pl)),
           campus_regions=us.groupby("gea").planned_it_mw.sum().div(1000).round(2).to_dict())
pt = stats("Mid-case")
boot = [stats("Mid-case", RNG.integers(0, len(w), len(w)), RNG.integers(0, len(pl), len(pl))) for _ in range(NBOOT)]
out["mid"] = pt
out["mid_ci95"] = {k: [float(np.nanpercentile([b[k] for b in boot], 2.5)), float(np.nanpercentile([b[k] for b in boot], 97.5))] for k in pt}
out["by_scenario"] = {s: stats(s) for s in SCEN}
states = us.name_1.to_numpy()
out["loo_avgscreen_change_mt"] = {s: stats("Mid-case", np.where(states != s)[0])["avgscreen_change_mt"] for s in sorted(set(states))}
json.dump(out, open("results/results_lrmer.json", "w"), indent=1)
print(json.dumps({k: out[k] for k in ("mid", "mid_ci95", "campus_regions")}, indent=1))
print({s: {k: round(v, 3) for k, v in d.items()} for s, d in out["by_scenario"].items()})
print("LOO avg-screen change:", min(out["loo_avgscreen_change_mt"].values()), max(out["loo_avgscreen_change_mt"].values()))
